import streamlit as st
import pandas as pd
import openpyxl
import io
import os
import base64      # <-- NEW IMPORT
import mimetypes   # <-- NEW IMPORT

# --- GLOBAL HELPER FUNCTIONS ---
def safe_write(sheet, r, c, val):
    """Safely writes a value to a cell, skipping merged or text-filled cells."""
    cell = sheet.cell(row=r, column=c)
    if type(cell).__name__ == 'MergedCell':
        return
    if isinstance(cell.value, str) and cell.value.strip() != "":
        return
    cell.value = val

def set_background(image_file):
    """Encodes an image/GIF and injects it as a CSS background."""
    # Guess the file type (e.g., image/gif, image/jpeg, image/png)
    mime_type, _ = mimetypes.guess_type(image_file)
    if mime_type is None:
        mime_type = "image/png" # Fallback
        
    try:
        with open(image_file, "rb") as f:
            data = f.read()
        encoded = base64.b64encode(data).decode()
        
        # CSS to inject the background into the main Streamlit container (.stApp)
        css = f"""
        <style>
        .stApp {{
            background-image: url("data:{mime_type};base64,{encoded}");
            background-size: cover;
            background-position: center;
            background-repeat: no-repeat;
            background-attachment: fixed;
        }}
        </style>
        """
        st.markdown(css, unsafe_allow_html=True)
    except FileNotFoundError:
        st.warning(f"پس منظر کی تصویر '{image_file}' نہیں ملی۔")

# --- APP SETUP ---
st.set_page_config(layout="wide", page_title="ٹرانزٹ آپریشنز رپورٹ جنریٹر")

# --- CALL THE BACKGROUND FUNCTION HERE ---
# Place your image or gif in the same folder as this script.
# Uncomment the line below and change "your_background.gif" to your actual file name.
# set_background("your_gif.gif") 

st.title("ٹرانزٹ آپریشنز رپورٹ جنریٹر")

# --- SIDEBAR NAVIGATION ---
st.sidebar.header("نیویگیشن")
app_mode = st.sidebar.radio(
    "چلانے کے لیے ٹول منتخب کریں:",
    ["اسٹیشن اے ایف سی رپورٹ", "کینٹ کارٹ ویلیڈیشن رپورٹ"]
)
st.sidebar.divider()

# ==========================================
# TOOL 1: STATION AFC REPORT
# ==========================================
if app_mode == "اسٹیشن اے ایف سی رپورٹ":
    st.subheader("گھنٹہ وار اسٹیشن ٹکٹ رپورٹ - ٹیمپلیٹ فلر")
    st.write("اپنا خام ٹکٹ ڈیٹا اپ لوڈ کریں۔ ایپ خود بخود ڈیفالٹ 'TAP Template.xlsx' استعمال کرے گی۔")

    st.sidebar.header("اے ایف سی رپورٹ کی ترتیبات")
    time_range = st.sidebar.slider(
        "رپورٹنگ کے اوقات منتخب کریں", 
        min_value=0, max_value=24, value=(6, 22), format="%d:00", key="afc_slider"
    )
    start_hr, end_hr = time_range

    # Keeping original names for code functionality
    TARGET_STATIONS = [
        "Faiz Ahmad Faiz", "G-13", "Golra More", 
        "N-5", "NHA", "NUST", "Police Foundation", "G-10"
    ]

    col1, col2 = st.columns(2)
    with col1:
        raw_file = st.file_uploader("1. خام ٹکٹ ڈیٹا اپ لوڈ کریں", type=["xlsx", "xls"], key="afc_raw")
    with col2:
        st.info("ڈیفالٹ ٹیمپلیٹ استعمال ہو رہا ہے: **TAP Template.xlsx**")
        template_file = st.file_uploader("اختیاری: ڈیفالٹ ٹیمپلیٹ تبدیل کریں", type=["xlsx"], key="afc_temp")

    # Determine which template to use (uploaded override vs. local default)
    default_template = "TAP Template.xlsx"
    active_template = template_file if template_file else (default_template if os.path.exists(default_template) else None)

    if raw_file is not None:
        if active_template is None:
            st.error(f"ایپ فولڈر میں ڈیفالٹ ٹیمپلیٹ '{default_template}' نہیں ملا۔ براہ کرم اسے دستی طور پر اپ لوڈ کریں۔")
        else:
            try:
                # 1. Process Raw Data
                df = pd.read_excel(raw_file, header=2)
                df['TIME'] = pd.to_datetime(df['TIME'])
                df['Hour_Int'] = df['TIME'].dt.hour
                
                def adjust_hour(row):
                    hr = row['Hour_Int']
                    station = row['STATION NAME']
                    if station in TARGET_STATIONS:
                        if hr < start_hr:
                            return start_hr                   
                        elif hr >= end_hr:
                            return max(start_hr, end_hr - 1)  
                    return hr
                    
                df['Adjusted_Hour'] = df.apply(adjust_hour, axis=1)
                df = df[(df['Adjusted_Hour'] >= start_hr) & (df['Adjusted_Hour'] < end_hr)]
                
                hourly_counts = df.groupby(['STATION NAME', 'Adjusted_Hour']).size().reset_index(name='Count')
                
                # 2. Process Template
                wb = openpyxl.load_workbook(active_template)
                ws = wb.active
                
                station_col_map = {}
                for col_idx in range(1, ws.max_column + 1):
                    cell_value = ws.cell(row=2, column=col_idx).value
                    if cell_value and isinstance(cell_value, str):
                        station_name = cell_value.strip()
                        for offset in range(3):
                            if ws.cell(row=3, column=col_idx + offset).value == "AFC":
                                station_col_map[station_name] = col_idx + offset
                                break
                                
                row_map = {h: h - 1 for h in range(start_hr, end_hr)}
                
                for st_name, col_idx in station_col_map.items():
                    for h, row_idx in row_map.items():
                        safe_write(ws, row_idx, col_idx, 0)
                        
                for _, row in hourly_counts.iterrows():
                    st_name = row['STATION NAME']
                    hr = row['Adjusted_Hour']
                    count = row['Count']
                    
                    if st_name in station_col_map and hr in row_map:
                        safe_write(ws, row_map[hr], station_col_map[st_name], count)
                
                # 3. Export
                buffer = io.BytesIO()
                wb.save(buffer)
                buffer.seek(0)
                
                st.success("اے ایف سی ٹیمپلیٹ کامیابی سے بھر گئی ہے!")
                with st.expander("نکالے گئے ڈیٹا کا پیش نظارہ (خام)"):
                    st.dataframe(hourly_counts)
                    
                st.download_button(
                    label="بھری ہوئی اے ایف سی رپورٹ ڈاؤن لوڈ کریں",
                    data=buffer,
                    file_name="Filled_AFC_Ridership_Report.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
                
            except Exception as e:
                st.error(f"ایک خامی پیش آ گئی: {e}")
                st.info("یقینی بنائیں کہ ٹیمپلیٹ کی ساخت متوقع فارمیٹ سے ملتی ہے۔")


# ==========================================
# TOOL 2: KENTKART VALIDATION REPORT
# ==========================================
elif app_mode == "کینٹ کارٹ ویلیڈیشن رپورٹ":
    st.subheader("کینٹ کارٹ ویلیڈیشن - ٹیمپلیٹ فلر")
    st.write("اپنا خام کینٹ کارٹ ڈیٹا اپ لوڈ کریں۔ ایپ خود بخود ڈیفالٹ 'KK Template.xlsx' استعمال کرے گی۔")

    st.sidebar.header("کینٹ کارٹ کی ترتیبات")
    time_range = st.sidebar.slider(
        "رپورٹنگ کے اوقات منتخب کریں", 
        min_value=0, max_value=24, value=(7, 18), format="%d:00", key="kentkart_slider"
    )
    start_hr, end_hr = time_range

    col1, col2 = st.columns(2)
    with col1:
        raw_file = st.file_uploader("1. خام کینٹ کارٹ ڈیٹا اپ لوڈ کریں", type=["xlsx", "xls"], key="kk_raw")
    with col2:
        st.info("ٹیمپلیٹ شامل کر دی گئی ہے: **KK Template.xlsx**")
        template_file = st.file_uploader("اختیاری: ڈیفالٹ ٹیمپلیٹ تبدیل کریں", type=["xlsx"], key="kk_temp")

    # Determine which template to use (uploaded override vs. local default)
    default_template = "KK Template.xlsx"
    active_template = template_file if template_file else (default_template if os.path.exists(default_template) else None)

    if raw_file is not None:
        if active_template is None:
            st.error(f"ایپ فولڈر میں ڈیفالٹ ٹیمپلیٹ '{default_template}' نہیں ملا۔ براہ کرم اسے دستی طور پر اپ لوڈ کریں۔")
        else:
            try:
                # 1. Process Raw Data
                df = pd.read_excel(raw_file, header=1)
                df.columns = df.columns.str.strip()
                
                target_col = 'Total Count'
                if target_col not in df.columns:
                    st.error(f"کالم '{target_col}' نہیں ملا۔ دستیاب کالم یہ ہیں: " + ", ".join(df.columns))
                    st.stop()
                    
                df[target_col] = pd.to_numeric(df[target_col], errors='coerce').fillna(0)
                df['Trip Start Date Time'] = pd.to_datetime(
                    df['Trip Start Date Time'].astype(str).str.strip(), 
                    errors='coerce'
                )
                df = df.dropna(subset=['Trip Start Date Time'])
                df['Hour_Int'] = df['Trip Start Date Time'].dt.hour
                
                def adjust_hour(hr):
                    if hr < start_hr:
                        return start_hr                   
                    elif hr >= end_hr:
                        return max(start_hr, end_hr - 1)  
                    return hr
                    
                df['Adjusted_Hour'] = df['Hour_Int'].apply(adjust_hour)
                
                def format_plate(p):
                    p = str(p).strip()
                    if p.startswith('EV') and '-' not in p:
                        return p.replace('EV', 'EV-')
                    return p
                    
                if 'Plate' in df.columns:
                    df['Formatted_Plate'] = df['Plate'].apply(format_plate)
                else:
                    st.error("کالم 'Plate' نہیں ملا۔")
                    st.stop()
                
                # --- NEW FEATURE: Get the last trip start time for each bus ---
                last_trip_times = df.groupby('Formatted_Plate')['Trip Start Date Time'].max().reset_index()
                last_trip_times.rename(columns={'Formatted_Plate': 'Bus Plate', 'Trip Start Date Time': 'Last Trip Start Time'}, inplace=True)
                
                # Filter data for hourly counts to ensure it stays strictly within the selected limits
                filtered_df = df[(df['Adjusted_Hour'] >= start_hr) & (df['Adjusted_Hour'] < end_hr)]
                hourly_counts = filtered_df.groupby(['Formatted_Plate', 'Adjusted_Hour'])[target_col].sum().reset_index(name='Count')
                unique_plates = sorted(filtered_df['Formatted_Plate'].dropna().unique())
                
                # 2. Process Template
                wb = openpyxl.load_workbook(active_template)
                ws = wb.active
                
                row_map = {h: h - 2 for h in range(start_hr, end_hr)}
                plate_col_map = {}
                
                for i, plate in enumerate(unique_plates):
                    # --- UPDATED: 4 columns per bus ---
                    # (e.g., Entered, Exited, Dispatch, Validation)
                    base_col = 2 + (i * 4)  
                    val_col = base_col + 3  # Validation is now the 4th column in the block
                    
                    plate_col_map[plate] = val_col
                    
                    # Assuming row 2 is still where the plate number goes
                    safe_write(ws, 2, base_col, plate)
                    
                    for hr, row_idx in row_map.items():
                        safe_write(ws, row_idx, val_col, 0)
                
                # ==========================================
                # ---> THIS IS THE MISSING BLOCK! <---
                # ==========================================
                for _, row in hourly_counts.iterrows():
                    plate = row['Formatted_Plate']
                    hr = row['Adjusted_Hour']
                    count = row['Count']
                    
                    if plate in plate_col_map and hr in row_map:
                        safe_write(ws, row_map[hr], plate_col_map[plate], count)
                # ==========================================
                        
                # 3. Export
                buffer = io.BytesIO()
                wb.save(buffer)
                buffer.seek(0)
                
                st.success("ویلیڈیشن ٹیمپلیٹ کامیابی سے بھر گئی ہے!")
                with st.expander("نکالے گئے ڈیٹا کا پیش نظارہ (ٹکٹوں کا مجموعہ)"):
                    st.dataframe(hourly_counts, use_container_width=True)
                    
                # --- DISPLAY NEW TABLE ---
                with st.expander("ہر بس کے آخری سفر کے آغاز کا وقت", expanded=True):
                    st.write("آپ اس ٹیبل کو کاپی کرنے کے لیے کلک اور ڈریگ کر سکتے ہیں، یا CSV کے طور پر ڈاؤن لوڈ کرنے کے لیے اس پر ہوور (hover) کر سکتے ہیں۔")
                    st.dataframe(last_trip_times, use_container_width=True)
                    
                st.download_button(
                    label="بھری ہوئی ویلیڈیشن رپورٹ ڈاؤن لوڈ کریں",
                    data=buffer,
                    file_name="Filled_Validation_Ridership_Report.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
                
            except KeyError as e:
                st.error(f"خام ڈیٹا میں متوقع کالم غائب ہے: {e}")
            except Exception as e:
                st.error(f"ایک خامی پیش آ گئی: {e}")
